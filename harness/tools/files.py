"""Atomic workspace-confined file operations (resolved paths reject symlinks out)."""
import asyncio
import os
from pathlib import Path
import tempfile


class FileTools:
    def __init__(self, workspace):
        self.workspace = Path(workspace).resolve()
        self.workspace.mkdir(parents=True, exist_ok=True)

    def resolve(self, path):
        target = (self.workspace / path).resolve()
        if not target.is_relative_to(self.workspace):
            raise PermissionError('Path is outside the configured workspace')
        return target

    async def read(self, path: str, max_chars: int = 50000):
        if max_chars < 0:
            raise ValueError('max_chars must be nonnegative')
        target = self.resolve(path)
        with target.open(encoding='utf-8') as stream:
            text = stream.read(max_chars + 1)
        return {'text': text[:max_chars], 'truncated': len(text) > max_chars}

    async def extract(self, path: str, max_chars: int = 50000):
        """Read PDF pages or workbook cells as data; no document macros are executed."""
        if not 1 <= max_chars <= 500000:
            raise ValueError('max_chars must be between 1 and 500000')
        target = self.resolve(path)
        def collect():
            if target.suffix.lower() == '.pdf':
                from pypdf import PdfReader
                parts = (f'Page {i+1}\n{page.extract_text() or ""}' for i, page in enumerate(PdfReader(target).pages))
            elif target.suffix.lower() == '.xlsx':
                from openpyxl import load_workbook
                workbook = load_workbook(target, read_only=True, data_only=True)
                try:
                    parts = (sheet.title + '\t' + '\t'.join('' if cell is None else str(cell) for cell in row)
                             for sheet in workbook for row in sheet.iter_rows(values_only=True))
                    return bounded(parts)
                finally:
                    workbook.close()
            else:
                raise ValueError('Extraction supports .pdf and .xlsx; use file_read for text or CSV')
            return bounded(parts)
        def bounded(parts):
            text = ''
            for part in parts:
                text += part[:max_chars+1] + '\n'
                if len(text) > max_chars:
                    return {'text': text[:max_chars], 'truncated': True}
            return {'text': text, 'truncated': False}
        return await asyncio.to_thread(collect)

    async def write(self, path: str, text: str):
        target = self.resolve(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        # Recheck after parent creation, before replacing.
        target = self.resolve(path)
        fd, name = tempfile.mkstemp(prefix='.write-', dir=target.parent)
        try:
            with os.fdopen(fd, 'w', encoding='utf-8', newline='') as stream:
                stream.write(text)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(name, target)
        finally:
            if os.path.exists(name):
                os.unlink(name)
        return {'path': target.relative_to(self.workspace).as_posix(), 'characters': len(text)}

    async def edit(self, path: str, old: str, new: str):
        if not old:
            raise ValueError('old must not be empty')
        target = self.resolve(path)
        text = target.read_text(encoding='utf-8')
        if text.count(old) != 1:
            raise ValueError('Edit requires exactly one occurrence of old text')
        return await self.write(path, text.replace(old, new, 1))

    async def list(self, path: str = '.'):
        target = self.resolve(path)
        return {'entries': [{'name': p.name, 'directory': p.is_dir()}
                            for p in sorted(target.iterdir())
                            if p.resolve().is_relative_to(self.workspace)]}
